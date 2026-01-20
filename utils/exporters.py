import os
from fpdf import FPDF
from docx import Document
from config import Config

def generate_srt(segments, output_path):
    """Generates SRT file from segments."""
    with open(output_path, 'w', encoding='utf-8') as f:
        for i, seg in enumerate(segments, start=1):
            start = format_timestamp(seg['start'])
            end = format_timestamp(seg['end'])
            text = seg['text'].strip()
            f.write(f"{i}\n{start} --> {end}\n{text}\n\n")
    return output_path

def generate_pdf(transcript_data, summary, output_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Transcription Report", ln=True, align='C')
    
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(200, 10, txt="Summary", ln=True, align='L')
    pdf.set_font("Arial", size=12)
    def clean_text(text):
        return text.encode('latin-1', 'replace').decode('latin-1')

    pdf.multi_cell(0, 10, txt=clean_text(summary))
    
    pdf.ln(10)
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(200, 10, txt="Transcript", ln=True, align='L')
    pdf.set_font("Arial", size=10)
    
    for seg in transcript_data:
        line = f"[{seg['start']:.2f}s] {seg.get('speaker', 'Unknown')}: {seg['text']}"
        pdf.multi_cell(0, 10, txt=clean_text(line))
        
    pdf.output(output_path)
    return output_path

def format_timestamp(seconds):
    """Helper for SRT timestamp formatting."""
    # Already defined in helpers.py but useful here if standalone.
    # We should ideally import from helpers to avoid duplication.
    millis = int((seconds - int(seconds)) * 1000)
    seconds = int(seconds)
    minutes, seconds = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours:02}:{minutes:02}:{seconds:02},{millis:03}"
